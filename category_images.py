"""Section (category) images for tl-shop.com.ua / Horoshop.

Usage:
    python category_images.py list "Біти та бітотримачі/Бітотримачі"   # products of a section
    python category_images.py build --test          # 10 test sections, both variants
    python category_images.py build --variant A     # all sections, one variant
    python category_images.py tree                  # print tree with product counts
    python category_images.py autopick              # fill missing leaf choices heuristically
    python category_images.py review --start 0      # contact sheet of chosen photos to check
    python category_images.py sheet ART [ART ...]   # contact sheet of any articles / sections

Input: export.xlsx.xlsx (Horoshop export). Choices: category_choices.csv
(path;article;reason). Output: category_images/{A,B}/cat-*.{jpg,png},
preview_{A,B}.png, category_images.xlsx.
"""
import argparse
import csv
import io
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd
import pymupdf
from openpyxl.styles import Alignment, Font
from PIL import Image, ImageChops, ImageDraw, ImageFont

import wera_photos as wp

ROOT = Path(__file__).resolve().parent
EXPORT = ROOT / "export.xlsx.xlsx"
CHOICES = ROOT / "category_choices.csv"
OUT = ROOT / "category_images"
CACHE = ROOT / "cache"
TARGET = 1200
MARGIN = 0.08
MIN_SOURCE = 600
JPEG_QUALITY = 90
REMBG_MODEL = "isnet-general-use"
WHITE_NOISE = 10   # source pixels this close to white count as background
WHITE_RAMP = 40    # ...and become fully opaque this far from white
HOLE_PURE_SHARE = 0.75  # enclosed spot is a real hole if this share of it is pure white

# sections present in the site menu but without products in the export
SITE_ONLY = [
    "Викрутки/Kraftform Plus – серія 300/Викрутки TORQ-SET® Mplus Kraftform Plus – серія 300",
    "Викрутки/Kraftform Plus – серія 300/Викрутки TRI-WING® Kraftform Plus – серія 300",
    "Діелектричні інструменти/Ізольовані викрутки Kraftform Plus – серія 100 VDE/"
    "Шестигранні ізольовані викрутки, Kraftform Plus – серія 100",
]

TEST = [
    "Біти та бітотримачі",
    "Біти та бітотримачі/Бітотримачі",
    "Біти та бітотримачі/Біти/Біти багатозубцеві",
    "Біти та бітотримачі/Біти/Біти для гвинтів ASSY®",
    "Викрутки",
    "Гайкові ключі Joker",
    "Сумки для інструментів Wera 2go",
    "Інструментальний візок та ложементи",
    "Діелектричні інструменти",
    "Тріскачки та приладдя Zyklop",
]


def log(msg):
    print(msg, flush=True)


# ---------- tree ----------

def split_path(p):
    # ¼ ⅜ ½ are single characters, so "/" is always a level separator
    return [s.strip() for s in str(p).split("/") if s.strip()]


def load_products():
    df = pd.read_excel(EXPORT, dtype=str).fillna("")
    df["Артикул"] = df["Артикул"].str.strip()
    return df


def build_tree(df):
    """Return (sections, own, extra): sections = {path: level}; own/extra = {path: [row idx]}."""
    sections, own, extra = {}, defaultdict(list), defaultdict(list)

    def add(path, idx, bucket):
        parts = split_path(path)
        for i in range(1, len(parts) + 1):
            sections.setdefault("/".join(parts[:i]), i)
        if parts:
            bucket["/".join(parts)].append(idx)

    for idx, row in df.iterrows():
        add(row["Раздел"], idx, own)
        for p in re.split(r"[;\n]", row["Дополнительные разделы"]):
            if p.strip():
                add(p, idx, extra)
    for p in SITE_ONLY:
        parts = split_path(p)
        for i in range(1, len(parts) + 1):
            sections.setdefault("/".join(parts[:i]), i)
    return sections, own, extra


def section_products(path, sections, own, extra):
    """Own products + products of all subsections + products listed via additional sections."""
    idxs = []
    for p in sections:
        if p == path or p.startswith(path + "/"):
            idxs += own.get(p, []) + extra.get(p, [])
    return list(dict.fromkeys(idxs))


# ---------- names ----------

TRANSLIT = {
    "а": "a", "б": "b", "в": "v", "г": "h", "ґ": "g", "д": "d", "е": "e", "є": "ie", "ж": "zh",
    "з": "z", "и": "y", "і": "i", "ї": "i", "й": "i", "к": "k", "л": "l", "м": "m", "н": "n",
    "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "kh", "ц": "ts",
    "ч": "ch", "ш": "sh", "щ": "shch", "ь": "", "ю": "iu", "я": "ia", "'": "", "’": "",
    "ы": "y", "э": "e", "ё": "e", "ъ": "",
    "¼": "1-4", "⅜": "3-8", "½": "1-2", "×": "x",
}


def translit(s):
    s = s.lower()
    s = "".join(TRANSLIT.get(c, c) for c in s)
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


UA_ABC = "абвгґдеєжзиіїйклмнопрстуфхцчшщьюя"


def path_key(path):
    """Sort sections in Ukrainian alphabetical order, level by level."""
    return [[UA_ABC.index(c) if c in UA_ABC else 100 + ord(c) for c in s.lower()]
            for s in split_path(path)]


def file_stem(path):
    return "cat-" + "--".join(translit(p) for p in split_path(path))


# ---------- sources ----------

def first_photo_url(row):
    urls = [u.strip() for u in re.split(r"[;\s]+", row["Фото"]) if u.strip().startswith("http")]
    return urls[0] if urls else None


def site_photo(url):
    name = re.sub(r"[^A-Za-z0-9._-]", "_", url.split("/content/images/")[-1])
    data = wp.cached(CACHE / "shop" / name, url, wp.is_image)
    if not data:
        return None
    im = Image.open(io.BytesIO(data))
    im.load()
    return im


def datasheet_photo(art):
    pdf = wp.fetch_pdf(art)
    if not pdf:
        return None
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    r = wp.pdf_product_image(doc)
    return r[0] if r else None


def source_image(row, prefer=None):
    """(Image, source label). Site photo first; datasheet if site photo is small/missing."""
    art = row["Артикул"]
    url = first_photo_url(row)
    if prefer != "pdf" and url:
        im = site_photo(url)
        if im is not None and max(im.size) >= MIN_SOURCE:
            return im, f"site {im.width}x{im.height}"
    im2 = datasheet_photo(art)
    if im2 is not None:
        return im2, f"wera.de datasheet {im2.width}x{im2.height}"
    if url:
        im = site_photo(url)
        if im is not None:
            return im, f"site {im.width}x{im.height} (small)"
    return None, "none"


# ---------- processing ----------

def has_alpha(im):
    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        a = im.convert("RGBA").getchannel("A")
        return a.getextrema()[0] < 250
    return False


def to_rgba(im):
    """Product on transparent background if the source already has one, else on white."""
    if has_alpha(im):
        return im.convert("RGBA")
    rgb = wp.to_rgb_white(im)
    return rgb.convert("RGBA")


def alpha_bbox(rgba):
    return rgba.getchannel("A").point(lambda v: 255 if v > 8 else 0).getbbox()


def white_bbox(rgb):
    return wp.content_bbox(rgb)


def place(prod, bg):
    """Fit product (RGBA) into TARGET square with MARGIN, keep proportions."""
    box = TARGET * (1 - 2 * MARGIN)
    scale = box / max(prod.size)
    new = (max(1, round(prod.width * scale)), max(1, round(prod.height * scale)))
    prod = prod.resize(new, Image.LANCZOS)
    canvas = Image.new("RGBA", (TARGET, TARGET), bg)
    canvas.alpha_composite(prod, ((TARGET - new[0]) // 2, (TARGET - new[1]) // 2))
    return canvas


_rembg_session = None


def remove_bg(rgba):
    global _rembg_session
    from rembg import new_session, remove
    if _rembg_session is None:
        _rembg_session = new_session(REMBG_MODEL)
    flat = Image.new("RGB", rgba.size, "white")
    flat.paste(rgba, mask=rgba.getchannel("A"))
    return remove(flat, session=_rembg_session).convert("RGBA")


def variant_a(im):
    rgba = to_rgba(im)
    flat = Image.new("RGB", rgba.size, "white")
    flat.paste(rgba, mask=rgba.getchannel("A"))
    prod = flat.crop(white_bbox(flat)).convert("RGBA")
    return place(prod, (255, 255, 255, 255)).convert("RGB")


def rembg_cut(im, cache_key):
    path = CACHE / "rembg" / f"{cache_key}.png"
    if path.exists():
        cut = Image.open(path)
        cut.load()
        return cut
    cut = remove_bg(to_rgba(im))
    path.parent.mkdir(parents=True, exist_ok=True)
    cut.save(path)
    return cut


def fill_highlights(a, rgb):
    """Make enclosed transparent spots opaque unless they are real see-through holes.
    A real hole (wrench ring, handle eye, bag strap) shows flat pure-white background;
    a blown-out chrome highlight is mostly a near-white gradient."""
    import numpy as np
    from scipy import ndimage
    solid = a > 0.5
    holes = ndimage.binary_fill_holes(solid) & ~solid
    lab, n = ndimage.label(holes)
    pure = rgb.min(axis=2) >= 252
    share = ndimage.mean(pure, lab, range(1, n + 1))
    real = np.isin(lab, [i + 1 for i, s in enumerate(share) if s >= HOLE_PURE_SHARE])
    # whole silhouette minus real holes is opaque; 2 px border keeps the soft edge
    body = ndimage.binary_erosion(ndimage.binary_fill_holes(solid) & ~real, iterations=2)
    a = a.copy()
    a[body] = 1
    return a


def cutout(im, cache_key):
    """Transparent cut-out. rembg loses thin chrome shafts on white photos, so its mask is
    merged with a 'not white' mask of the source; edge colours are un-blended from white."""
    import numpy as np
    if has_alpha(im):
        return im.convert("RGBA")  # source already has a real transparent background
    rgb = np.asarray(wp.to_rgb_white(im), dtype=np.float32)
    a_rembg = np.asarray(rembg_cut(im, cache_key).getchannel("A"), dtype=np.float32) / 255
    diff = (255 - rgb).max(axis=2)
    a_white = np.clip((diff - WHITE_NOISE) / WHITE_RAMP, 0, 1)
    a = np.maximum(a_rembg, a_white)
    a = fill_highlights(a, rgb)
    safe = np.maximum(a, 1e-3)[..., None]
    col = np.clip((rgb - (1 - a[..., None]) * 255) / safe, 0, 255)
    out = np.dstack([col, a * 255]).round().astype(np.uint8)
    return Image.fromarray(out, "RGBA")


def variant_b(im, cache_key):
    cut = cutout(im, cache_key)
    bbox = alpha_bbox(cut) or (0, 0, cut.width, cut.height)
    return place(cut.crop(bbox), (0, 0, 0, 0))


# ---------- preview ----------

def font(size):
    for f in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
              "/usr/share/fonts/dejavu/DejaVuSans.ttf"):
        if Path(f).exists():
            return ImageFont.truetype(f, size)
    return ImageFont.load_default()


def wrap(draw, text, fnt, width):
    lines, cur = [], ""
    for w in text.split():
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=fnt) <= width:
            cur = t
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def preview(items, path, cols=5, tile=260):
    """items: [(name, Image)]. Dark site-like grid: picture + section name under it."""
    bg, card, fg = (0x1b, 0x1a, 0x19), (0x26, 0x25, 0x24), (0xf2, 0xf2, 0xf2)
    pad, text_h = 24, 70
    rows = (len(items) + cols - 1) // cols
    W = pad + cols * (tile + pad)
    H = pad + rows * (tile + text_h + pad)
    sheet = Image.new("RGB", (W, H), bg)
    d = ImageDraw.Draw(sheet)
    fnt = font(15)
    for i, (name, im) in enumerate(items):
        x = pad + (i % cols) * (tile + pad)
        y = pad + (i // cols) * (tile + text_h + pad)
        d.rounded_rectangle([x, y, x + tile, y + tile + text_h], 10, fill=card)
        th = im.convert("RGBA").resize((tile - 20, tile - 20), Image.LANCZOS)
        sheet.paste(th, (x + 10, y + 10), th)
        for j, line in enumerate(wrap(d, name, fnt, tile - 16)[:3]):
            tw = d.textlength(line, font=fnt)
            d.text((x + (tile - tw) / 2, y + tile + 4 + j * 20), line, font=fnt, fill=fg)
    sheet.save(path)


# ---------- choices / table ----------

def read_choices():
    if not CHOICES.exists():
        return {}
    with open(CHOICES, encoding="utf-8") as f:
        return {r["path"]: r for r in csv.DictReader(f, delimiter=";")}


def write_choices(choices):
    with open(CHOICES, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["path", "article", "source", "reason"], delimiter=";")
        w.writeheader()
        for p in sorted(choices, key=path_key):
            w.writerow({k: choices[p].get(k, "") for k in w.fieldnames})


BAG_WORDS = ("неукомплект", "без інструмент", "без інструмента", "запасн")
BAG_START = ("сумка", "м'яка сумка", "м'який футляр", "футляр", "складна сумка", "ложемент",
             "кожух", "текстильна", "поясна кобура", "сумка-скручування", "смужка")


def is_bag(name):
    n = name.lower().lstrip("0123456789 ")
    return any(w in n for w in BAG_WORDS) or n.startswith(BAG_START)


def is_set(name):
    n = name.lower()
    return "набір" in n or re.search(r"\bset\b", n) is not None or "ложемент" in n


def is_set_section(name):
    n = name.lower()
    return "набор" in n or "набір" in n or "набори" in n


def auto_pick(path, df, sections, own, extra, used):
    """Leaf heuristic: topical own products, sets vs single tools by section name, median size."""
    name = split_path(path)[-1]
    own_idx = own.get(path, [])
    pool = own_idx or section_products(path, sections, own, extra)
    names = {i: df.loc[i, "Название (UA)"] for i in pool}
    cand = [i for i in pool if not is_bag(names[i])] or pool
    want_set = is_set_section(name)
    pref = [i for i in cand if is_set(names[i]) == want_set] or cand
    fresh = [i for i in pref if df.loc[i, "Артикул"] not in used] or pref
    i = fresh[len(fresh) // 2]
    kind = "набор — раздел про наборы" if want_set and is_set(names[i]) else \
        "штучный инструмент — средний типоразмер/вариант раздела" if not is_set(names[i]) else \
        "набор (штучных товаров в разделе нет)"
    return df.loc[i, "Артикул"], kind


def cmd_autopick(df, sections, own, extra, opts):
    choices = read_choices()
    used = {c["article"] for c in choices.values()}
    leaves = [p for p in sections if not any(q.startswith(p + "/") for q in sections)]
    n = 0
    for p in sorted(leaves):
        if p in choices or not section_products(p, sections, own, extra):
            continue
        art, why = auto_pick(p, df, sections, own, extra, used)
        used.add(art)
        choices[p] = {"path": p, "article": art, "source": "", "reason": "auto: " + why}
        n += 1
    write_choices(choices)
    log(f"autopick: {n} new, {len(choices)} total")


def write_xlsx(rows, empty, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Разделы"
    head = ["Полный путь раздела", "Название раздела", "Уровень", "Файл",
            "Артикул-источник", "Название товара", "Почему выбран"]
    ws.append(head)
    for r in rows:
        ws.append([r[h] for h in head])
    widths = [60, 40, 9, 50, 16, 60, 60]
    for ws_ in [ws]:
        for i, w in enumerate(widths):
            ws_.column_dimensions[chr(65 + i)].width = w
    ws2 = wb.create_sheet("Без товаров")
    ws2.append(["Полный путь раздела", "Название раздела", "Уровень", "Примечание"])
    for r in empty:
        ws2.append(r)
    ws2.column_dimensions["A"].width = 90
    ws2.column_dimensions["B"].width = 50
    ws2.column_dimensions["D"].width = 50
    for w in (ws, ws2):
        for c in w[1]:
            c.font = Font(bold=True)
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=False, vertical="top")
        for row in w.iter_rows(min_row=2, min_col=5, max_col=5):
            for c in row:
                c.number_format = "@"
        w.freeze_panes = "A2"
    wb.save(path)


# ---------- commands ----------

def cmd_tree(df, sections, own, extra, opts):
    for p in sorted(sections):
        n = len(section_products(p, sections, own, extra))
        print(f"{'  ' * (sections[p] - 1)}{split_path(p)[-1]}  [{n}]")
    print(len(sections), "sections")


def cmd_list(df, sections, own, extra, opts):
    for path in opts.paths:
        idxs = section_products(path, sections, own, extra)
        print(f"== {path}: {len(idxs)} products")
        for i in idxs:
            r = df.loc[i]
            sub = r["Раздел"][len(path) + 1:] if r["Раздел"].startswith(path) else "+" + r["Раздел"]
            print(f"{r['Артикул']} | {r['Название (UA)'][:110]} | {sub[:70]}")


def cmd_sheet(df, sections, own, extra, opts):
    """Contact sheet of source photos for visual check: articles or 'section path'."""
    by_art = {r["Артикул"]: i for i, r in df.iterrows()}
    arts = []
    for a in opts.items:
        if a in by_art:
            arts.append(a)
        else:
            arts += [df.loc[i, "Артикул"] for i in section_products(a, sections, own, extra)]
    items = []
    for a in arts[:opts.limit]:
        row = df.loc[by_art[a]]
        im, src = source_image(row, opts.source)
        if im is None:
            im = Image.new("RGB", (100, 100), "red")
        items.append((f"{a} {src}", variant_a(im)))
    preview(items, Path(opts.out), cols=opts.cols, tile=300)
    log(f"sheet -> {opts.out} ({len(items)})")


def cmd_review(df, sections, own, extra, opts):
    """Sheet of chosen source photos: '#n section | article' for visual check."""
    choices = read_choices()
    by_art = {r["Артикул"]: i for i, r in df.iterrows()}
    paths = sorted(choices, key=path_key)
    items = []
    for n, p in enumerate(paths[opts.start:opts.start + opts.count], opts.start):
        ch = choices[p]
        im, src = source_image(df.loc[by_art[ch["article"]]], ch.get("source") or None)
        im = variant_a(im) if im is not None else Image.new("RGB", (100, 100), "red")
        items.append((f"#{n} {split_path(p)[-1][:60]} | {ch['article']} {src.split()[0]}", im))
    preview(items, Path(opts.out), cols=6, tile=280)
    log(f"review -> {opts.out}: {opts.start}..{opts.start + len(items) - 1} of {len(paths)}")


def cmd_build(df, sections, own, extra, opts):
    choices = read_choices()
    paths = TEST if opts.test else sorted(sections, key=path_key)
    variants = ["A", "B"] if opts.variant == "AB" else [opts.variant]
    by_art = {r["Артикул"]: i for i, r in df.iterrows()}
    rows, empty = [], []
    prev = {v: [] for v in variants}
    for n, path in enumerate(paths, 1):
        name, level = split_path(path)[-1], sections.get(path, len(split_path(path)))
        idxs = section_products(path, sections, own, extra)
        ch = choices.get(path)
        if not idxs or not ch:
            if path not in sections:
                note = "раздела нет ни в выгрузке, ни в меню сайта"
            elif not idxs:
                note = "нет товаров в выгрузке (раздел есть в меню сайта)"
            else:
                note = "нет выбора в category_choices.csv"
            empty.append([path, name, level, note])
            log(f"[{n}/{len(paths)}] {path}: SKIP ({note})")
            continue
        art = ch["article"]
        row = df.loc[by_art[art]]
        im, src = source_image(row, ch.get("source") or None)
        log(f"[{n}/{len(paths)}] {path} <- {art} ({src})")
        if im is None:
            empty.append([path, name, level, f"нет фото у {art}"])
            continue
        stem = file_stem(path)
        files = []
        for v in variants:
            d = OUT / v
            d.mkdir(parents=True, exist_ok=True)
            if v == "A":
                out = variant_a(im)
                fn = stem + ".jpg"
                out.save(d / fn, "JPEG", quality=JPEG_QUALITY, optimize=True)
            else:
                out = variant_b(im, f"{art}-{ch.get('source') or 'auto'}")
                fn = stem + ".png"
                out.save(d / fn, "PNG", optimize=True)
            files.append(fn)
            prev[v].append((name, out))
        rows.append({"Полный путь раздела": path, "Название раздела": name, "Уровень": level,
                     "Файл": " | ".join(files), "Артикул-источник": art,
                     "Название товара": row["Название (UA)"],
                     "Почему выбран": ch["reason"] + (f" [источник: {src}]" if src else "")})
    suffix = "_test" if opts.test else ""
    for v in variants:
        preview(prev[v], ROOT / f"preview_{v}{suffix}.png", cols=5 if opts.test else 8)
    write_xlsx(rows, empty, ROOT / f"category_images{suffix}.xlsx")
    log(f"done: {len(rows)} sections with images, {len(empty)} without")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("tree")
    p = sub.add_parser("list")
    p.add_argument("paths", nargs="+")
    p = sub.add_parser("sheet")
    p.add_argument("items", nargs="+", help="articles or section paths")
    p.add_argument("--out", default="sheet.png")
    p.add_argument("--cols", type=int, default=5)
    p.add_argument("--limit", type=int, default=40)
    p.add_argument("--source", default=None, choices=[None, "pdf"])
    p = sub.add_parser("build")
    p.add_argument("--test", action="store_true")
    p.add_argument("--variant", default="AB", choices=["A", "B", "AB"])
    sub.add_parser("autopick")
    p = sub.add_parser("review")
    p.add_argument("--start", type=int, default=0)
    p.add_argument("--count", type=int, default=30)
    p.add_argument("--out", default="review.png")
    opts = ap.parse_args()
    df = load_products()
    sections, own, extra = build_tree(df)
    cmds = {"tree": cmd_tree, "list": cmd_list, "sheet": cmd_sheet, "build": cmd_build,
            "autopick": cmd_autopick, "review": cmd_review}
    cmds[opts.cmd](df, sections, own, extra, opts)


if __name__ == "__main__":
    main()
